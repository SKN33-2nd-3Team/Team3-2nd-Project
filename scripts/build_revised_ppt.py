# -*- coding: utf-8 -*-
"""Rebuild PlaylistPro final presentation into required 14-slide storyline,
preserving the existing design system. Reuses 9 existing slides, builds 5 new."""
from pptx import Presentation
from pptx.util import Inches as In, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from PIL import Image
import copy, os

SRC="outputs/PlaylistPro_final_presentation_project_centered.pptx"
OUT="outputs/PlaylistPro_final_presentation_revised.pptx"
FONT="Malgun Gothic"

# palette
NAVY="17354B"; NAVY2="21465F"; INK="17212B"; CORAL="E85E43"; CORAL_L="F6A28F"
MINT="18847A"; MINT_BG="EDF7F4"; CORAL_BG="FFF3EF"; GREY_BG="F3F6F8"
GREY="60707E"; LINE="D9E1E6"; WHITE="FFFFFF"; SKY="D9E5EC"; SKYD="AFC3D0"

def C(h): return RGBColor.from_string(h)

prs=Presentation(SRC)
layout=prs.slide_masters[0].slide_layouts[0]

def strip_placeholders(slide):
    for ph in list(slide.placeholders):
        ph._element.getparent().remove(ph._element)

def set_bg(slide,hexcol):
    cSld=slide._element.find(qn('p:cSld'))
    # remove existing bg
    old=cSld.find(qn('p:bg'))
    if old is not None: cSld.remove(old)
    bg=cSld.makeelement(qn('p:bg'),{})
    bgPr=bg.makeelement(qn('p:bgPr'),{})
    sf=bgPr.makeelement(qn('a:solidFill'),{})
    clr=sf.makeelement(qn('a:srgbClr'),{'val':hexcol})
    sf.append(clr); bgPr.append(sf)
    fillref=bgPr.makeelement(qn('a:effectLst'),{}); bgPr.append(fillref)
    bg.append(bgPr)
    cSld.insert(0,bg)

def no_line(sh): sh.line.fill.background()
def no_shadow(sh):
    sh.shadow.inherit=False

def box(slide,x,y,w,h,paras,anchor=MSO_ANCHOR.TOP):
    """paras: list of paragraphs; each paragraph = list of run tuples (text,size,color,bold,align).
       For simple single-format paragraph pass (text,size,color,bold,align)."""
    tb=slide.shapes.add_textbox(In(x),In(y),In(w),In(h)); tf=tb.text_frame
    tf.word_wrap=True; tf.vertical_anchor=anchor
    tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0
    first=True
    for pa in paras:
        p=tf.paragraphs[0] if first else tf.add_paragraph()
        first=False
        runs = pa if isinstance(pa,list) else [pa]
        align=runs[0][4] if len(runs[0])>4 else PP_ALIGN.LEFT
        p.alignment=align
        for rt in runs:
            text,size,color,bold=rt[0],rt[1],rt[2],rt[3]
            r=p.add_run(); r.text=text
            r.font.name=FONT; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=C(color)
    return tb

def rrect(slide,x,y,w,h,fill,line=None,radius=0.045):
    sh=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,In(x),In(y),In(w),In(h))
    sh.fill.solid(); sh.fill.fore_color.rgb=C(fill)
    if line: sh.line.color.rgb=C(line); sh.line.width=Pt(0.75)
    else: no_line(sh)
    no_shadow(sh)
    try: sh.adjustments[0]=radius
    except: pass
    return sh

def rect(slide,x,y,w,h,fill,line=None):
    sh=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,In(x),In(y),In(w),In(h))
    sh.fill.solid(); sh.fill.fore_color.rgb=C(fill)
    if line: sh.line.color.rgb=C(line); sh.line.width=Pt(0.75)
    else: no_line(sh)
    no_shadow(sh); return sh

def framed_pic(slide,path,fx,fy,fw,fh,pad=0.12,frame=True):
    if frame:
        rrect(slide,fx,fy,fw,fh,WHITE,line=LINE,radius=0.03)
    iw,ih=Image.open(path).size; ar=iw/ih
    availw=fw-2*pad; availh=fh-2*pad
    if availw/availh > ar:
        ph=availh; pw=ph*ar
    else:
        pw=availw; ph=pw/ar
    px=fx+(fw-pw)/2; py=fy+(fh-ph)/2
    slide.shapes.add_picture(path,In(px),In(py),In(pw),In(ph))

def common_footer(slide,srcnote):
    rect(slide,0.58,7.04,12.17,0.012,LINE)
    box(slide,0.58,7.15,5.83,0.17,[("PlaylistPro · Retention Decision Support",7.5,GREY,False,PP_ALIGN.LEFT)])
    if srcnote:
        box(slide,0.58,6.78,11.56,0.18,[("근거: "+srcnote,7.5,GREY,False,PP_ALIGN.LEFT)])

def header(slide,eyebrow,title,subtitle,title_size=28.5,eyecolor=CORAL):
    box(slide,0.58,0.17,6.5,0.23,[(eyebrow,9,eyecolor,True,PP_ALIGN.LEFT)])
    box(slide,0.58,0.44,12.16,1.10,[(title,title_size,INK,True,PP_ALIGN.LEFT)])
    box(slide,0.58,1.59,11.77,0.30,[(subtitle,12,GREY,False,PP_ALIGN.LEFT)])

def add_pagenum(slide,n,navy=False):
    col=SKYD if navy else GREY
    y=0.40 if navy else 7.12
    box(slide,12.10,y,0.65,0.19,[(f"{n:02d}",8.25,col,True,PP_ALIGN.RIGHT)])

def new_slide(navy=False):
    s=prs.slides.add_slide(layout)
    strip_placeholders(s)
    set_bg(s, NAVY if navy else WHITE)
    return s

FIG="figures/presentation_v3/"
EDA="artifacts/eda_insight/"
SS="submission_package/PlaylistPro_final_submission_20260722/assets/screenshots/"

# ============ NEW SLIDE 7: 8개 후보 모델 비교 ============
s7=new_slide()
header(s7,"MODEL BENCHMARK",
       "8개 모델을 동일한 조건에서 비교해 상위 후보를 선별했습니다",
       "같은 log_numeric 피처·고정 5-Fold·동일 OOF로 8개 모델을 한 번에 비교했다")
framed_pic(s7,FIG+"04_eight_model_fair_comparison.png",0.58,2.00,7.75,4.10)
# right ranked list card
rrect(s7,8.55,2.00,4.20,4.10,GREY_BG,radius=0.05)
box(s7,8.82,2.22,3.6,0.28,[("5-Fold OOF PR-AUC",11.5,NAVY,True,PP_ALIGN.LEFT)])
ranks=[("CatBoost","0.9476",True),("LightGBM","0.9475",True),("XGBoost","0.9471",True),
       ("Gradient Boosting","0.9464",False),("Random Forest","0.9394",False),
       ("Decision Tree","0.9361",False),("Logistic","0.9063",False),("Dummy","0.5134",False)]
yy=2.60
for name,val,top in ranks:
    col=MINT if top else INK
    box(s7,8.82,yy,2.55,0.24,[(("★ " if top else "· ")+name,11,col,top,PP_ALIGN.LEFT)])
    box(s7,11.35,yy,1.20,0.24,[(val,11,col,top,PP_ALIGN.RIGHT)])
    yy+=0.335
box(s7,8.82,5.42,3.6,0.55,[("★ 상위 3개(CatBoost·XGBoost·LightGBM)를 정밀 튜닝·안정성 검증 대상으로 선별",10,GREY,False,PP_ALIGN.LEFT)])
add_pagenum(s7,7)
common_footer(s7,"04_eight_model_fair_comparison.csv · model_comparison_fair.csv · run_manifest.json")
s7.notes_slide.notes_text_frame.text=(
"이 슬라이드의 결론은 8개 모델을 완전히 같은 조건에서 한 번에 비교했다는 것입니다. "
"모두 동일한 log_numeric 피처와 고정 5-Fold, 같은 OOF 예측으로 평가했기 때문에 점수 차이는 데이터나 폴드가 아니라 모델의 차이입니다. "
"Dummy는 0.5134로 하한선이고, Logistic 0.9063, 트리·부스팅 계열이 0.936에서 0.948 사이입니다. "
"상위 세 개인 CatBoost, LightGBM, XGBoost는 소수점 넷째 자리에서 갈릴 만큼 근소합니다. "
"주의할 점은 이 표가 최종 선택이 아니라 후보 선별이라는 것입니다. 다음 장에서 이 상승이 어떤 단계에서 만들어졌는지 봅니다.")

# ============ NEW SLIDE 9: 최종 모델 선정 ============
s9=new_slide()
header(s9,"MODEL SELECTION",
       "CatBoost는 근소한 순위 우위로 선택했지만 절대적인 승자는 아닙니다",
       "상위 3개 정밀 튜닝·5-Seed·Bootstrap으로 검증 · Recall 우선이면 LightGBM도 합리적 대안",
       title_size=26)
framed_pic(s9,FIG+"06_top3_fine_tuning_before_after.png",0.58,2.00,5.95,3.75)
framed_pic(s9,FIG+"08_bootstrap_confidence_intervals.png",6.68,2.00,6.07,3.75)
band=rrect(s9,0.58,5.90,12.17,0.66,GREY_BG,radius=0.10)
box(s9,0.85,6.00,11.6,0.24,[[
    ("OOF PR-AUC  ",10.5,NAVY,True),
    ("CatBoost 0.947897 · LightGBM 0.947790 · XGBoost 0.947765",10.5,INK,False)]])
box(s9,0.85,6.28,11.6,0.24,[[
    ("안정성  ",10.5,NAVY,True),
    ("5-Seed 평균 CatBoost 0.947880 (σ=0.000061) · Bootstrap 95% CI 중첩   |   ",10.5,INK,False),
    ("Trade-off  ",10.5,CORAL,True),
    ("Recall@0.50 LightGBM 0.8599 > CatBoost 0.8364 · Brier LightGBM 근소 우수",10.5,INK,False)]])
add_pagenum(s9,9)
common_footer(s9,"final_model_selection_matrix.csv · top3_fine_tuning_summary.csv · 07_seed_stability.csv · 08_bootstrap_confidence_intervals.csv")
s9.notes_slide.notes_text_frame.text=(
"결론부터 말하면 CatBoost는 조건부 기술 후보이지 압도적 승자가 아닙니다. "
"최종 OOF PR-AUC는 CatBoost 0.947897, LightGBM 0.947790, XGBoost 0.947765로 CatBoost가 근소하게 앞섭니다. "
"5개 Seed 평균도 CatBoost가 가장 높고 표준편차 0.000061로 변동이 작습니다. "
"다만 Bootstrap 신뢰구간이 서로 겹치고, 0.5 임계값 기준 Recall과 Brier는 LightGBM이 더 좋습니다. "
"그래서 순위 품질과 반복 안정성을 우선해 CatBoost를 골랐지만, 놓침을 최소화하려면 LightGBM도 합리적이라고 정직하게 설명합니다. "
"모델 선택과 임계값 선택은 별개라는 점이 다음 장으로 이어집니다. 먼저 이 모델을 어떻게 저장했는지 봅니다.")

# ============ NEW SLIDE 10: 모델 저장과 재로딩 검증 ============
s10=new_slide()
header(s10,"MODEL PACKAGING",
       "최종 모델은 재현 가능한 Pipeline과 Metadata로 저장했습니다",
       "전처리와 CatBoost를 하나의 Pipeline으로 고정 · Streamlit이 동일 저장 모델을 재로딩한다")
# left card: saved artifacts
rrect(s10,0.58,2.02,6.02,3.55,MINT_BG,radius=0.05)
box(s10,0.85,2.22,5.5,0.28,[("저장 산출물",13,MINT,True,PP_ALIGN.LEFT)])
left_rows=[
 ("모델 파일","music_churn_pipeline.joblib"),
 ("저장 위치","artifacts/model/"),
 ("구성","전처리 + CatBoost 결합 Pipeline"),
 ("Feature schema","log_numeric · 19개 입력 컬럼"),
 ("최종 모델","CatBoost"),
 ("OOF PR-AUC","0.947897"),
 ("Run ID","20260721_full_fair_v1"),
]
yy=2.62
for k,v in left_rows:
    box(s10,0.85,yy,1.85,0.26,[(k,10.5,GREY,False,PP_ALIGN.LEFT)])
    box(s10,2.75,yy,3.6,0.26,[(v,10.5,INK,True,PP_ALIGN.LEFT)])
    yy+=0.40
# right card: integrity + reload
rrect(s10,6.73,2.02,6.02,3.55,GREY_BG,radius=0.05)
box(s10,7.00,2.22,5.5,0.28,[("무결성·재로딩 검증",13,NAVY,True,PP_ALIGN.LEFT)])
box(s10,7.00,2.62,1.85,0.26,[("SHA-256",10.5,GREY,False,PP_ALIGN.LEFT)])
box(s10,8.90,2.62,3.75,0.26,[("fc35ca91e924 … 989a7d0e",10.5,INK,True,PP_ALIGN.LEFT)])
checks=[
 ("파일 SHA 일치","music_churn_pipeline.joblib 검증"),
 ("reload_verified","true"),
 ("새 프로세스 재로딩","true (fresh_process)"),
 ("소스 후보","models/candidates/…/candidate_pipeline.joblib"),
 ("로드 규칙","저장소 내부 경로 + SHA-256 검증 후 로드"),
]
yy=3.02
for k,v in checks:
    box(s10,7.00,yy,2.05,0.26,[(k,10.5,GREY,False,PP_ALIGN.LEFT)])
    box(s10,9.05,yy,3.6,0.26,[(v,10.5,INK,True,PP_ALIGN.LEFT)])
    yy+=0.40
# bottom callout band
band=rrect(s10,0.58,5.78,12.17,0.78,NAVY,radius=0.09)
box(s10,0.85,5.90,1.0,0.28,[("의미",11,CORAL_L,True,PP_ALIGN.LEFT)])
box(s10,0.85,6.20,11.9,0.28,[("학습 결과를 Streamlit이 동일하게 재사용하도록 모델·전처리·Metadata를 하나의 Pipeline으로 함께 고정했습니다.",13.5,WHITE,True,PP_ALIGN.LEFT)])
add_pagenum(s10,10)
common_footer(s10,"artifacts/model/metadata.json · run_manifest.json · artifacts/model/music_churn_pipeline.joblib")
s10.notes_slide.notes_text_frame.text=(
"이 슬라이드의 결론은 모델을 파일로 저장한 데서 끝내지 않고 재현 가능하게 고정했다는 것입니다. "
"저장 산출물은 artifacts/model 폴더의 music_churn_pipeline.joblib 하나이며, 전처리와 CatBoost가 결합된 단일 Pipeline입니다. "
"메타데이터에는 최종 모델 CatBoost, OOF PR-AUC 0.947897, Run ID 20260721_full_fair_v1가 기록돼 있습니다. "
"무결성은 SHA-256 fc35ca91e9242dbc0e914db9f0dfc9e994024a5d4cc8d0d770d12321989a7d0e로 검증했고, "
"저장 후 같은 프로세스뿐 아니라 새 프로세스에서도 재로딩을 확인했습니다(reload_verified·fresh_process_reload_verified 모두 true). "
"핵심 의미는 Streamlit이 학습을 다시 하지 않고 검증된 이 저장 모델만 로드한다는 것입니다. 다음 장에서 그 앱을 봅니다.")

# ============ NEW SLIDE 11: Streamlit 시연 1 — 운영 기준 선택 ============
s11=new_slide()
header(s11,"STREAMLIT · OPERATING CONTROLS",
       "Streamlit에서 Threshold와 Top-K를 실제 검토 인원으로 바꿉니다",
       "정답 하나가 아니라 자원과 오류 비용에 맞춰 운영자가 선택한다 · Lift는 위험 농축도")
framed_pic(s11,FIG+"10_threshold_precision_recall_f1.png",0.58,1.98,6.02,3.05)
framed_pic(s11,FIG+"14_topk_capture_lift.png",6.73,1.98,6.02,3.05)
# four scenario mini boxes
scen=[
 (MINT_BG,MINT,"재현율 우선 · Th 0.29","77,700명 · Recall 95.18%","FN 3,091 · FP 16,617"),
 (GREY_BG,NAVY,"균형형 F1 · Th 0.35","76,143명 · Recall 94.44%","FN 3,569 · FP 15,538"),
 (CORAL_BG,CORAL,"정밀도 우선 · Th 0.74","48,463명 · Precision 95.07%","FN 18,099 · FP 2,388"),
 (GREY_BG,NAVY,"Top 10% 용량","12,500명 · Capture 19.48%","Lift 1.95 (농축도)"),
]
bx=0.58; bw=2.94; gap=0.11
for i,(bg,acc,t,l1,l2) in enumerate(scen):
    x=bx+i*(bw+gap)
    rrect(s11,x,5.18,bw,1.02,bg,radius=0.06)
    box(s11,x+0.18,5.30,bw-0.32,0.26,[(t,11,acc,True,PP_ALIGN.LEFT)])
    box(s11,x+0.18,5.64,bw-0.32,0.24,[(l1,10.5,INK,False,PP_ALIGN.LEFT)])
    box(s11,x+0.18,5.90,bw-0.32,0.24,[(l2,10.5,GREY,False,PP_ALIGN.LEFT)])
box(s11,0.58,6.33,12.17,0.26,[("Lift는 캠페인 uplift가 아니라 위험 고객 농축도입니다. 0.35는 OOF F1 기본값일 뿐 사업 최적값이 아닙니다.",11,INK,True,PP_ALIGN.LEFT)])
add_pagenum(s11,11)
common_footer(s11,"10_11_threshold_sweep.csv · 14_topk_capture_lift.csv · artifacts/model/metadata.json")
s11.notes_slide.notes_text_frame.text=(
"이 슬라이드의 결론은 하나의 정답 임계값이 없다는 것입니다. Streamlit에서 운영자가 직접 손잡이를 돌립니다. "
"임계값 0.29는 7만 7,700명을 검토해 Recall 95.18%를 얻지만 FP가 1만 6,617명입니다. "
"0.35는 비슷하게 넓은 균형안이고, 0.74는 대상이 4만 8,463명으로 줄고 정밀도 95.07% 대신 FN이 1만 8,099명으로 늘어납니다. "
"인원이 먼저 정해져 있으면 Top-K가 더 직관적입니다. 상위 10% 1만 2,500명은 이탈자의 19.48%를 포착하고 Lift는 1.95입니다. "
"여기서 반드시 강조할 한계는 Lift 1.95가 캠페인으로 이탈이 1.95배 준다는 뜻이 아니라, 위험 고객이 무작위 대비 얼마나 농축됐는지라는 것입니다. "
"이렇게 고른 고객을 다음 화면에서 신호와 함께 검토합니다.")

# ============ NEW SLIDE 14: 한계·다음 검증·최종 결론 (navy) ============
s14=new_slide(navy=True)
box(s14,0.58,0.40,6.5,0.25,[("PROJECT OUTCOME",9,CORAL_L,True,PP_ALIGN.LEFT)])
box(s14,0.58,0.82,11.7,0.80,[("모델 점수를 실행 가능한 고객 검토 순서로 바꿨습니다",30,WHITE,True,PP_ALIGN.LEFT)])
box(s14,0.58,1.66,11.4,0.40,[("성과는 점수 하나가 아니라 고객 인사이트·모델 검증·운영 기준·유지 행동을 연결한 의사결정 구조입니다.",13.5,SKY,False,PP_ALIGN.LEFT)])
steps=[("01","고객 이해","낮은 청취·Free·높은 문의 신호 발견"),
       ("02","공정한 검증","고정 5-Fold OOF·8개 모델 비교"),
       ("03","운영 선택","Threshold·Top-K로 검토 범위 결정"),
       ("04","행동 검토","투명한 규칙과 담당자 최종 선택")]
sx=0.58; sw=2.94; sgap=0.11
for i,(n,t,b) in enumerate(steps):
    x=sx+i*(sw+sgap)
    accent=CORAL if i==3 else MINT
    rect(s14,x,2.42,sw,0.03,CORAL if i==3 else "6E8797")
    box(s14,x,2.62,1.0,0.29,[(n,12,CORAL_L,True,PP_ALIGN.LEFT)])
    box(s14,x,3.06,sw,0.32,[(t,16.5,WHITE,True,PP_ALIGN.LEFT)])
    box(s14,x,3.52,sw-0.1,0.60,[(b,11.25,SKY,False,PP_ALIGN.LEFT)])
# limitations strip
rrect(s14,0.58,4.42,12.17,0.86,NAVY2,radius=0.07)
box(s14,0.85,4.53,11.9,0.24,[("반드시 함께 말할 한계",10.5,CORAL_L,True,PP_ALIGN.LEFT)])
box(s14,0.85,4.84,11.9,0.34,[("실제 미래 라벨 Holdout 없음 · 예측 기간 미정 · 규칙 기반 합성 데이터 · 실제 이탈 감소 미검증 · 실제 ROI 미검증 · 행동 효과는 A/B 테스트 필요",11.5,SKYD,False,PP_ALIGN.LEFT)])
# conclusion band
rrect(s14,0.58,5.46,12.17,1.10,NAVY2,radius=0.06)
box(s14,0.85,5.62,1.0,0.28,[("결론",12,CORAL_L,True,PP_ALIGN.LEFT)])
box(s14,0.85,5.96,11.9,0.55,[("PlaylistPro의 성과는 PR-AUC 숫자 하나가 아니라, 고객 인사이트·모델 검증·운영 기준·유지 행동을 하나의 의사결정 흐름으로 연결한 것입니다.",15,WHITE,True,PP_ALIGN.LEFT)])
add_pagenum(s14,14,navy=True)
s14.notes_slide.notes_text_frame.text=(
"마지막 슬라이드의 결론은, 이 프로젝트가 이탈 여부를 맞히는 데서 멈추지 않고 누구부터 왜 볼지를 정하는 구조를 만들었다는 것입니다. "
"고객 이해에서 위험 신호를 찾고, 고정 5-Fold OOF로 8개 모델을 공정하게 비교했으며, Threshold와 Top-K로 검토 범위를 정하고, 투명한 규칙으로 행동을 담당자가 선택하게 했습니다. "
"핵심 수치는 CatBoost OOF PR-AUC 0.947897이지만, 그 숫자 하나가 성과가 아닙니다. "
"과장하면 안 되는 한계도 분명합니다. 실제 미래 라벨 Holdout이 없고, 예측 기간이 정의되지 않았으며, 합성 데이터로 판단했고, 실제 이탈 감소와 ROI는 검증되지 않았습니다. "
"그래서 마지막 문장은 이것입니다. 성과는 점수가 아니라 인사이트·검증·운영·행동을 하나의 의사결정 흐름으로 연결한 것입니다. 여기서 발표를 마치고 Q&A로 넘어갑니다.")

# ---- update images/text on kept slide 9(idx8 original MODEL COMPARISON) : we DELETE it, not keep ----

# ============ REORDER + DELETE ============
ex=list(prs.slides)  # current order: 0..18 (14 original + 5 new appended)
# original indices
o=ex[:14]
# desired order of slide objects
desired=[o[0],o[1],o[2],o[3],   # 1-4
         o[5],o[6],              # 5 feature(ex6), 6 validation(ex7)
         s7,                     # 7
         o[7],                   # 8 performance(ex8)
         s9,                     # 9
         s10,                    # 10
         s11,                    # 11
         o[11],                  # 12 product(ex12)
         o[4],                   # 13 retention(ex5)
         s14]                    # 14
delete_set=[o[8],o[9],o[10],o[12],o[13]]  # ex9,ex10,ex11,ex13,ex14

# build part->rId map
part_to_rid={}
for rId,rel in prs.part.rels.items():
    if rel.reltype.endswith('/slide') and not rel.is_external:
        part_to_rid[rel.target_part]=rId
sldIdLst=prs.slides._sldIdLst
rid_to_el={sldId.get(qn('r:id')):sldId for sldId in list(sldIdLst)}
# drop rels + sldId for deleted
for sl in delete_set:
    rid=part_to_rid[sl.part]
    el=rid_to_el[rid]; sldIdLst.remove(el)
    prs.part.drop_rel(rid)
# reorder remaining to desired
# remove all remaining then append in order
for el in list(sldIdLst): sldIdLst.remove(el)
for sl in desired:
    rid=part_to_rid[sl.part]
    sldIdLst.append(rid_to_el[rid])

# ============ RENUMBER page numbers on kept slides ============
# kept slides that already have a page-number textbox (bottom-right ~x12.29). Update text to final position.
final_order=desired
def set_kept_pagenum(slide,n):
    # find shape near x>=11.9 with short numeric text
    target=None
    for sh in slide.shapes:
        if sh.has_text_frame and sh.left is not None and Emu(sh.left).inches>=11.8 and Emu(sh.width).inches<1.2:
            t=sh.text_frame.text.strip()
            if t.isdigit():
                target=sh; break
    if target is not None:
        for p in target.text_frame.paragraphs:
            for r in p.runs:
                r.text=f"{n:02d}"
        # clear extra runs
        return True
    return False

# slide1 (cover) has NO page number -> add one
cover=final_order[0]
add_pagenum(cover,1)  # bottom-right on white area
for i,sl in enumerate(final_order):
    n=i+1
    if sl in (s7,s9,s10,s11,s14,cover): continue  # already numbered
    set_kept_pagenum(sl,n)

prs.save(OUT)
print("SAVED",OUT,"slides=",len(prs.slides._sldIdLst))
